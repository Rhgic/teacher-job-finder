"""从上传的简历文件里抽取纯文本。

只做"拿到文字"这一步，栏目切分交给 resume_parser.structure_resume_text，
两者都不调用 LLM——上传阶段不该产生模型费用。

刻意不支持的格式：
- .doc（Word 97-2003 二进制）：需要 antiword/libreoffice 这类外部程序，
  为一个演示项目引入系统依赖不划算；提示用户另存为 .docx。
- 图片（jpg/png）：需要 OCR 引擎（tesseract + 中文语言包），
  同样是系统依赖，且中文简历识别质量不稳定。
两者都在下方 UNSUPPORTED_HINTS 里给出明确指引，而不是笼统报"格式不支持"。
"""
from __future__ import annotations

import io

# 单文件上限。简历极少超过 1MB，留 5MB 余量同时防止大文件占满内存。
MAX_UPLOAD_BYTES = 5 * 1024 * 1024

TEXT_SUFFIXES = {".txt", ".md", ".markdown", ".text"}
SUPPORTED_SUFFIXES = TEXT_SUFFIXES | {".pdf", ".docx"}

UNSUPPORTED_HINTS = {
    ".doc": "旧版 .doc 无法直接解析，请用 Word 另存为 .docx 后再上传",
    ".jpg": "暂不支持图片简历（需 OCR），请上传 PDF/Word，或直接粘贴文本",
    ".jpeg": "暂不支持图片简历（需 OCR），请上传 PDF/Word，或直接粘贴文本",
    ".png": "暂不支持图片简历（需 OCR），请上传 PDF/Word，或直接粘贴文本",
    ".webp": "暂不支持图片简历（需 OCR），请上传 PDF/Word，或直接粘贴文本",
    ".pages": "Pages 文档请先导出为 PDF 或 Word 再上传",
}


class ExtractError(ValueError):
    """抽取失败，message 直接面向用户展示。"""


def suffix_of(filename: str) -> str:
    name = (filename or "").strip().lower()
    return name[name.rfind("."):] if "." in name else ""


def extract_text(filename: str, data: bytes) -> str:
    """按扩展名分派抽取。失败抛 ExtractError，消息可直接给用户看。"""
    if len(data) > MAX_UPLOAD_BYTES:
        raise ExtractError(f"文件超过 {MAX_UPLOAD_BYTES // 1024 // 1024}MB 上限")
    if not data:
        raise ExtractError("文件是空的")

    suffix = suffix_of(filename)
    if suffix in UNSUPPORTED_HINTS:
        raise ExtractError(UNSUPPORTED_HINTS[suffix])
    if suffix not in SUPPORTED_SUFFIXES:
        raise ExtractError("只支持 PDF、Word(.docx)、Markdown、txt")

    if suffix in TEXT_SUFFIXES:
        text = _decode_text(data)
    elif suffix == ".pdf":
        text = _extract_pdf(data)
    else:
        text = _extract_docx(data)

    text = text.strip()
    if not text:
        raise ExtractError("没能从文件里读出文字——可能是扫描件或纯图片排版，请改用粘贴")
    return text


def _decode_text(data: bytes) -> str:
    # 简历常见于 UTF-8；国内 Windows 导出的 txt 多为 GBK，兜一层再放弃。
    for encoding in ("utf-8", "gb18030"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ExtractError("文本编码无法识别，请另存为 UTF-8")


def _extract_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(data))
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    except ExtractError:
        raise
    except Exception as exc:  # noqa: BLE001 - 第三方解析异常统一转为用户可读信息
        raise ExtractError(f"PDF 解析失败：{exc}") from exc


def _extract_docx(data: bytes) -> str:
    from docx import Document

    try:
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        # 不少简历把内容排在表格里，只读段落会得到空文本
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    parts.append(" ".join(cells))
        return "\n".join(parts)
    except ExtractError:
        raise
    except Exception as exc:  # noqa: BLE001 - 同上
        raise ExtractError(f"Word 解析失败：{exc}") from exc
