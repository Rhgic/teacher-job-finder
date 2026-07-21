"""简历文件文本抽取测试。

重点不只是"能抽出文字"，还包括拒绝路径给出的提示是否可操作——
用户上传 .doc 或截图时，需要知道下一步该做什么，而不是笼统的"格式不支持"。
"""
import io

import pytest

from services.resume_extract import MAX_UPLOAD_BYTES, ExtractError, extract_text


def test_markdown_and_txt_are_read_as_text():
    data = "# 陈晓\n小学体育教师资格证".encode()
    assert "小学体育教师资格证" in extract_text("resume.md", data)
    assert "小学体育教师资格证" in extract_text("resume.txt", data)


def test_gbk_encoded_text_falls_back():
    """国内 Windows 导出的 txt 常是 GBK，不该直接失败。"""
    data = "华南师范大学 体育教育".encode("gb18030")
    assert "体育教育" in extract_text("resume.txt", data)


def test_docx_reads_paragraphs_and_tables():
    """不少简历把内容排在表格里，只读段落会得到空文本。"""
    from docx import Document

    doc = Document()
    doc.add_paragraph("陈晓 体育教师")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "证书"
    table.rows[0].cells[1].text = "国家二级运动员"
    buf = io.BytesIO()
    doc.save(buf)

    text = extract_text("resume.docx", buf.getvalue())
    assert "陈晓 体育教师" in text
    assert "国家二级运动员" in text


def test_oversized_file_is_rejected():
    with pytest.raises(ExtractError, match="上限"):
        extract_text("resume.pdf", b"x" * (MAX_UPLOAD_BYTES + 1))


def test_empty_file_is_rejected():
    with pytest.raises(ExtractError, match="空"):
        extract_text("resume.md", b"")


def test_image_hint_points_to_a_way_out():
    """图片简历要告诉用户改用什么，而不是只说不支持。"""
    with pytest.raises(ExtractError) as exc:
        extract_text("resume.jpg", b"\xff\xd8\xff fake jpeg")
    assert "粘贴" in str(exc.value)


def test_legacy_doc_hint_points_to_a_way_out():
    with pytest.raises(ExtractError) as exc:
        extract_text("resume.doc", b"fake doc")
    assert ".docx" in str(exc.value)


def test_unknown_suffix_lists_supported_formats():
    with pytest.raises(ExtractError) as exc:
        extract_text("resume.xlsx", b"fake")
    assert "PDF" in str(exc.value)


def test_textless_pdf_tells_user_to_paste():
    """扫描件抽不出文字时，应引导改用粘贴而不是留下空简历。"""
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    buf = io.BytesIO()
    writer.write(buf)

    with pytest.raises(ExtractError, match="粘贴"):
        extract_text("scan.pdf", buf.getvalue())
